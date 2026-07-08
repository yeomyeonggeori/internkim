#!/usr/bin/env python3
import argparse
import re
from pathlib import Path

from skill_runtime import ensure_requirements

INLINE_PATTERN = re.compile(r"(\*\*.+?\*\*|\*.+?\*|`.+?`)")


def parse_arguments():
    parser = argparse.ArgumentParser(description="Render a markdown source of truth into a .docx deliverable")
    parser.add_argument("markdown_path", help="path to content.md")
    parser.add_argument("--output", help="output .docx path; defaults next to the markdown")
    parser.add_argument("--font", default="맑은 고딕", help="base font family name")
    parser.add_argument("--font-size", type=float, default=10.5)
    return parser.parse_args()


def add_inline_runs(paragraph, text, point_class):
    for segment in INLINE_PATTERN.split(text):
        if not segment:
            continue
        if segment.startswith("**") and segment.endswith("**") and len(segment) > 4:
            run = paragraph.add_run(segment[2:-2])
            run.bold = True
        elif segment.startswith("*") and segment.endswith("*") and len(segment) > 2:
            run = paragraph.add_run(segment[1:-1])
            run.italic = True
        elif segment.startswith("`") and segment.endswith("`") and len(segment) > 2:
            run = paragraph.add_run(segment[1:-1])
            run.font.name = "Courier New"
            run.font.size = point_class(9.5)
        else:
            paragraph.add_run(segment)


def is_table_line(line):
    stripped = line.strip()
    return stripped.startswith("|") and stripped.endswith("|") and stripped.count("|") >= 2


def is_divider_row(line):
    return bool(re.fullmatch(r"\|?[\s:|-]+\|?", line.strip())) and "-" in line


def split_table_row(line):
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def render_table(document, table_lines, point_class):
    rows = [split_table_row(line) for line in table_lines if not is_divider_row(line)]
    if not rows:
        return
    column_count = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=column_count)
    table.style = "Table Grid"
    for row_index, row in enumerate(rows):
        for column_index in range(column_count):
            cell = table.rows[row_index].cells[column_index]
            text = row[column_index] if column_index < len(row) else ""
            cell.paragraphs[0].text = ""
            add_inline_runs(cell.paragraphs[0], text, point_class)
            if row_index == 0:
                for run in cell.paragraphs[0].runs:
                    run.bold = True


def set_base_font(document, font_name, font_size, point_class):
    from docx.oxml.ns import qn

    style = document.styles["Normal"]
    style.font.name = font_name
    style.font.size = point_class(font_size)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), font_name)


def render_markdown(document, markdown_text, point_class):
    lines = markdown_text.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if not stripped:
            index += 1
            continue
        heading_match = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if heading_match:
            document.add_heading(heading_match.group(2).strip(), level=len(heading_match.group(1)))
            index += 1
            continue
        if is_table_line(line):
            table_lines = []
            while index < len(lines) and is_table_line(lines[index]):
                table_lines.append(lines[index])
                index += 1
            render_table(document, table_lines, point_class)
            continue
        bullet_match = re.match(r"^\s*[-*]\s+(.*)$", line)
        if bullet_match:
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline_runs(paragraph, bullet_match.group(1), point_class)
            index += 1
            continue
        numbered_match = re.match(r"^\s*\d+[.)]\s+(.*)$", line)
        if numbered_match:
            paragraph = document.add_paragraph(style="List Number")
            add_inline_runs(paragraph, numbered_match.group(1), point_class)
            index += 1
            continue
        if stripped.startswith(">"):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = point_class(18)
            run_text = stripped.lstrip("> ").strip()
            add_inline_runs(paragraph, run_text, point_class)
            for run in paragraph.runs:
                run.italic = True
            index += 1
            continue
        paragraph_lines = []
        while index < len(lines) and lines[index].strip() and not re.match(r"^(#{1,4})\s+", lines[index].strip()) and not is_table_line(lines[index]) and not re.match(r"^\s*([-*]|\d+[.)])\s+", lines[index]):
            paragraph_lines.append(lines[index].strip())
            index += 1
        paragraph = document.add_paragraph()
        add_inline_runs(paragraph, " ".join(paragraph_lines), point_class)


def main():
    arguments = parse_arguments()
    ensure_requirements("document")
    from docx import Document
    from docx.shared import Pt

    markdown_path = Path(arguments.markdown_path)
    markdown_text = markdown_path.read_text(encoding="utf-8")
    output_path = Path(arguments.output) if arguments.output else markdown_path.with_suffix(".docx")

    document = Document()
    set_base_font(document, arguments.font, arguments.font_size, Pt)
    render_markdown(document, markdown_text, Pt)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    print(f"exported {output_path} from {markdown_path}")


if __name__ == "__main__":
    main()
