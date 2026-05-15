#!/usr/bin/env python3
import argparse
import csv
import json
from pathlib import Path

from skill_runtime import ensure_requirements


def require_text(value, field_name):
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f"{field_name} must be a non-empty string")
    return value.strip()


def load_specification(specification_path):
    with open(specification_path, "r", encoding="utf-8") as specification_file:
        specification = json.load(specification_file)
    if not isinstance(specification, dict):
        raise ValueError("workbook specification must be an object")
    return specification


def create_workbook(specification):
    if not ensure_requirements("xlsx"):
        raise RuntimeError("xlsx dependencies are unavailable after bootstrap")

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)

    sheets = specification.get("sheets", [])
    if not isinstance(sheets, list) or not sheets:
        raise ValueError("sheets must be a non-empty array")

    for sheet_specification in sheets:
        worksheet = add_sheet(workbook, sheet_specification)
        apply_default_formatting(worksheet, Font, PatternFill, get_column_letter)

    return workbook


def add_sheet(workbook, sheet_specification):
    if not isinstance(sheet_specification, dict):
        raise ValueError("each sheet must be an object")
    title = require_text(sheet_specification.get("title"), "sheet.title")
    worksheet = workbook.create_sheet(title=title[:31])
    rows = read_rows(sheet_specification)
    for row in rows:
        worksheet.append(["" if value is None else value for value in row])
    freeze_panes = sheet_specification.get("freezePanes")
    if isinstance(freeze_panes, str) and freeze_panes.strip():
        worksheet.freeze_panes = freeze_panes.strip()
    if sheet_specification.get("autoFilter") and worksheet.max_row > 0 and worksheet.max_column > 0:
        worksheet.auto_filter.ref = worksheet.dimensions
    return worksheet


def read_rows(sheet_specification):
    if "csvPath" in sheet_specification:
        return read_delimited_rows(sheet_specification)
    rows = sheet_specification.get("rows", [])
    if not isinstance(rows, list):
        raise ValueError("sheet rows must be an array")
    for row in rows:
        if not isinstance(row, list):
            raise ValueError("each row must be an array")
    return rows


def read_delimited_rows(sheet_specification):
    csv_path = require_text(sheet_specification.get("csvPath"), "csvPath")
    delimiter = sheet_specification.get("delimiter", ",")
    if delimiter == "\\t":
        delimiter = "\t"
    with open(csv_path, newline="", encoding="utf-8-sig") as delimited_file:
        return list(csv.reader(delimited_file, delimiter=delimiter))


def apply_default_formatting(worksheet, font_class, fill_class, get_column_letter):
    if worksheet.max_row == 0:
        return
    header_fill = fill_class("solid", fgColor="E5E7EB")
    for cell in worksheet[1]:
        cell.font = font_class(bold=True)
        cell.fill = header_fill
    for column_cells in worksheet.columns:
        content_width = max(len(str(cell.value or "")) for cell in column_cells)
        column_letter = get_column_letter(column_cells[0].column)
        worksheet.column_dimensions[column_letter].width = min(max(content_width + 2, 10), 48)


def parse_arguments():
    parser = argparse.ArgumentParser(description="Create an XLSX workbook from a JSON specification.")
    parser.add_argument("specification_path")
    parser.add_argument("output_path")
    return parser.parse_args()


def main():
    arguments = parse_arguments()
    specification = load_specification(arguments.specification_path)
    workbook = create_workbook(specification)
    output_path = Path(arguments.output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    print(output_path)


if __name__ == "__main__":
    main()
