#!/usr/bin/env python3
import argparse
import json

from skill_runtime import ensure_requirements


def summarize_workbook(workbook_path):
    if not ensure_requirements("xlsx"):
        raise RuntimeError("xlsx dependencies are unavailable after bootstrap")

    from openpyxl import load_workbook

    workbook = load_workbook(workbook_path, data_only=False)
    sheets = []
    formula_cells = []
    for worksheet in workbook.worksheets:
        sheets.append({
            "title": worksheet.title,
            "rows": worksheet.max_row,
            "columns": worksheet.max_column,
        })
        for row in worksheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formula_cells.append(f"{worksheet.title}!{cell.coordinate}")
    return {
        "sheetCount": len(workbook.worksheets),
        "sheets": sheets,
        "formulaCells": formula_cells[:50],
        "formulaCellCount": len(formula_cells),
    }


def parse_arguments():
    parser = argparse.ArgumentParser(description="Validate and summarize an XLSX workbook.")
    parser.add_argument("workbook_path")
    return parser.parse_args()


def main():
    arguments = parse_arguments()
    summary = summarize_workbook(arguments.workbook_path)
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
